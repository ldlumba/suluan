import io
import json
import tempfile
import unittest
import zipfile
from email.message import Message
from unittest.mock import patch

import pymupdf
from docx import Document

from extraction import ExtractionError, extract
from pilot_fixtures import build_pilot_fixtures
from cloud_runner import (
    ServiceWorkerClient,
    decode_jwt_payload,
    format_worker_error,
    metadata_identity_token,
    process_claim,
    run_claims,
)
from runner import bounded_extract, WorkerClient, run_job
from operator_client import OperatorWorkerClient
from semantic_prepare import run_batches

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ExtractionTests(unittest.TestCase):
    def pdf(self, text=""):
        with pymupdf.open() as doc:
            page = doc.new_page()
            if text:
                page.insert_text((72, 72), text)
            return doc.tobytes()

    def test_pdf_text_and_provenance(self):
        result = extract(self.pdf("ILAW study findings"), PDF)
        self.assertIn("ILAW study findings", result["text"])
        self.assertEqual(result["pageCount"], 1)
        self.assertEqual(len(result["sha256"]), 64)
        self.assertEqual(result["warnings"], [])

    def test_image_only_or_blank_pdf_needs_attention(self):
        with self.assertRaisesRegex(ExtractionError, "NO_EXTRACTABLE_TEXT"):
            extract(self.pdf(), PDF)

    def test_mixed_pdf_warns_of_missing_text(self):
        with pymupdf.open(stream=self.pdf("Text page"), filetype="pdf") as doc:
            doc.new_page()
            result = extract(doc.tobytes(), PDF)
        self.assertEqual(result["warnings"], ["PAGES_WITHOUT_TEXT"])

    def test_docx_includes_tables_in_order(self):
        doc = Document()
        doc.add_paragraph("Before")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table finding"
        doc.add_paragraph("After")
        data = io.BytesIO()
        doc.save(data)
        result = extract(data.getvalue(), DOCX)
        self.assertEqual(result["text"], "Before\n\nTable finding\n\nAfter")
        self.assertIsNone(result["pageCount"])

    def test_malformed_files(self):
        for mime in (PDF, DOCX):
            with self.assertRaisesRegex(ExtractionError, "INVALID_DOCUMENT"):
                extract(b"not a document", mime)

    def test_encrypted_pdf(self):
        with pymupdf.open(stream=self.pdf("secret"), filetype="pdf") as doc:
            data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
        with self.assertRaisesRegex(ExtractionError, "ENCRYPTED_DOCUMENT"):
            extract(data, PDF)

    def test_stage_five_generated_failure_fixtures(self):
        expected = {
            "F01": (PDF, "NO_EXTRACTABLE_TEXT"),
            "F02": (PDF, "INVALID_DOCUMENT"),
            "F03": (PDF, "ENCRYPTED_DOCUMENT"),
            "F04": (DOCX, "INVALID_DOCUMENT"),
        }
        with tempfile.TemporaryDirectory() as directory:
            fixtures = build_pilot_fixtures(directory)
            self.assertEqual(set(fixtures), set(expected))
            for fixture_id, path in fixtures.items():
                mime_type, error_code = expected[fixture_id]
                with self.subTest(fixture_id=fixture_id):
                    with self.assertRaisesRegex(ExtractionError, error_code):
                        extract(path.read_bytes(), mime_type)

    def test_zip_entry_limit_before_parsing(self):
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w") as archive:
            for i in range(2001):
                archive.writestr(str(i), "x")
        with self.assertRaisesRegex(ExtractionError, "PROCESSING_LIMIT"):
            extract(data.getvalue(), DOCX)

    def test_subprocess_runner(self):
        self.assertEqual(bounded_extract(self.pdf("Local test"), PDF)["status"], "EXTRACTED")

    def test_token_cannot_be_sent_to_remote_url(self):
        with self.assertRaises(ValueError):
            WorkerClient("https://example.com", "test-token")
        with self.assertRaises(ValueError):
            OperatorWorkerClient("https://example.com", "test-token")
        with self.assertRaises(ValueError):
            OperatorWorkerClient(
                "http://example.com", "test-token", allow_remote=True
            )
        with self.assertRaises(ValueError):
            ServiceWorkerClient("http://example.com", lambda: "test-token")

    def test_semantic_operator_requires_explicit_https_remote_opt_in(self):
        client = OperatorWorkerClient(
            "https://worker.example", "test-token", allow_remote=True
        )
        self.assertEqual(client.base, "https://worker.example")
        self.assertEqual(client.token, "test-token")

    def test_runner_downloads_claimed_file_and_posts_fenced_result(self):
        data = self.pdf("Integration fixture")
        calls = []

        class Client:
            def json(self, path, body):
                calls.append((path, body))
                return {"submissionId": "accepted-id", "claimToken": "attempt-token"}

            def request(self, path):
                self_path = "/api/submissions/accepted-id/download"
                if path != self_path:
                    raise AssertionError(path)
                response = io.BytesIO(data)
                response.headers = Message()
                response.headers["Content-Type"] = PDF
                return response

        status, error = run_job(Client(), "job-id")
        self.assertEqual((status, error), ("EXTRACTED", None))
        payload = calls[-1][1]
        self.assertEqual(payload["claimToken"], "attempt-token")
        self.assertIn("Integration fixture", payload["result"]["text"])

    def test_explicit_processing_limits(self):
        with patch("extraction.MAX_TEXT", 5):
            with self.assertRaisesRegex(ExtractionError, "PROCESSING_LIMIT"):
                extract(self.pdf("Too much extracted text"), PDF)
        self.assertEqual(bounded_extract(self.pdf("Timeout"), PDF, timeout=0)["errorCode"], "PROCESSING_LIMIT")

    def test_semantic_operator_paces_successful_batches(self):
        calls, waits, output = [], [], []

        class Client:
            results = iter(
                [
                    {"embedded": 20, "papers": 1, "hasMore": True},
                    {"embedded": 3, "papers": 1, "hasMore": False},
                ]
            )

            def json(self, path, body):
                calls.append((path, body))
                return next(self.results)

        total = run_batches(
            Client(), interval=30, sleep=waits.append, emit=output.append
        )
        self.assertEqual(total, 23)
        self.assertEqual(
            calls,
            [
                ("/api/search/semantic/prepare", {}),
                ("/api/search/semantic/prepare", {}),
            ],
        )
        self.assertEqual(waits, [30])
        self.assertEqual(output[-1], "semantic preparation complete")

    def test_cloud_runner_pulls_bounded_claims_and_posts_result(self):
        data = self.pdf("Cloud extraction")
        calls = []

        class Response(io.BytesIO):
            def __init__(self):
                super().__init__(data)
                self.headers = Message()
                self.headers["Content-Type"] = PDF

            def __enter__(self):
                return self

            def __exit__(self, *_):
                self.close()

        class Client:
            claims = iter(
                [
                    {"jobId": "job-id", "claimToken": "claim-token"},
                    None,
                ]
            )

            def json(self, path, body=None):
                calls.append((path, body))
                if path.endswith("/claim"):
                    return next(self.claims)
                return {"status": body["status"]}

            def request(self, path):
                calls.append((path, None))
                return Response()

        output = []
        self.assertEqual(run_claims(Client(), emit=output.append), 1)
        self.assertIn("claimToken=claim-token", calls[1][0])
        self.assertEqual(calls[2][1]["status"], "EXTRACTED")
        self.assertEqual(output[-1], "no claimable jobs; completed=1")

    def test_cloud_runner_requests_an_audience_bound_metadata_token(self):
        captured = {}

        class Opener:
            def open(self, request, timeout):
                captured["url"] = request.full_url
                captured["header"] = request.get_header("Metadata-flavor")
                captured["timeout"] = timeout
                return io.BytesIO(b"signed-token\n")

        token = metadata_identity_token("https://worker.example", Opener())
        self.assertEqual(token, "signed-token")
        self.assertIn("audience=https%3A%2F%2Fworker.example", captured["url"])
        self.assertIn("format=full", captured["url"])
        self.assertEqual(captured["header"], "Google")
        self.assertEqual(captured["timeout"], 10)


    def test_cloud_runner_sends_browser_compatible_user_agent(self):
        import urllib.error

        captured = {}

        class FakeResponse:
            status = 204

            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

        class Opener:
            def open(self, request, timeout):
                captured["ua"] = request.get_header("User-agent")
                return FakeResponse()

        client = ServiceWorkerClient("https://worker.example", lambda: "token")
        client.opener = Opener()
        self.assertIsNone(client.json("/api/internal/processing/claim", {}))
        self.assertIn("SULUAN-Processing-Job", captured["ua"])

    def test_cloud_runner_error_reports_path_body_and_token_claims(self):
        import base64
        import urllib.error

        payload = base64.urlsafe_b64encode(
            json.dumps(
                {
                    "aud": "https://worker.example",
                    "email": "svc@example.iam.gserviceaccount.com",
                }
            ).encode()
        ).decode().rstrip("=")
        token = f"header.{payload}.signature"
        body = io.BytesIO(b'{"error":"forbidden"}')
        error = urllib.error.HTTPError(
            "https://worker.example/api/internal/processing/claim",
            403,
            "Forbidden",
            {},
            body,
        )
        detail = format_worker_error(error, token)
        self.assertIn("/api/internal/processing/claim", detail)
        self.assertIn("403", detail)
        self.assertIn("forbidden", detail)
        self.assertIn("https://worker.example", detail)
        self.assertIn("svc@example.iam.gserviceaccount.com", detail)
        self.assertNotIn(token, detail)
        self.assertEqual(decode_jwt_payload("not-a-jwt"), {})

    def test_cloud_runner_download_failure_posts_failed_without_retry(self):
        import urllib.error

        posted = []

        class Client:
            last_token = ""

            def json(self, path, body=None):
                posted.append((path, body))
                return {"status": body["status"]}

            def request(self, path):
                raise urllib.error.HTTPError(path, 403, "Forbidden", {}, io.BytesIO(b"{}"))

        status, error = process_claim(
            Client(), {"jobId": "job-id", "claimToken": "claim-token"}
        )
        self.assertEqual((status, error), ("FAILED", "DOWNLOAD_FAILED"))
        self.assertEqual(posted[0][1]["errorCode"], "DOWNLOAD_FAILED")


if __name__ == "__main__":
    unittest.main()
